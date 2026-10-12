"""Transcribe Table S1 of Lombardo, Berellini and Obach (2018) from the supplementary workbook to CSV.

Usage: python3 -I make_lombardo2018_csv.py WORKBOOK.xlsx OUTPUT.csv WORKBOOK_EXPORT.csv

WORKBOOK_EXPORT.csv is the comma-separated export of the same workbook that sits beside it in the repository; it is
read only to quote, in the `note` column, how a long-decimal cell is displayed there.

Nothing is computed or estimated. Each numeric cell is written with the shortest decimal text that reads back to
the stored number; a text cell in a numeric column is left empty and quoted in the `note` column.
"""
import csv
import sys

import openpyxl

HEADER_ROW = 9
SOURCE_HEADER = ["Name", "CAS #", "SMILES", "human VDss (L/kg)", "human CL (mL/min/kg)",
                 "fraction unbound \nin plasma (fu)", "MRT (h)", "terminal  t1/2 (h)", "Reference", "Comments ", "Notes",
                 "Year of first disclosure", "MW", "HBA", "HBD", "TPSA_NO", "RotBondCount", "moka_ionState7.4",
                 "MoKa.LogP", "MoKa.LogD7.4"]
OUTPUT_HEADER = ["name", "cas_number", "smiles", "vdss_l_per_kg", "clearance_ml_per_min_per_kg",
                 "fraction_unbound_plasma", "mrt_h", "half_life_h", "reference", "source_comments", "source_notes",
                 "year_of_first_disclosure", "mw", "hba", "hbd", "tpsa_no", "rot_bond_count", "moka_ion_state_7_4",
                 "moka_logp", "moka_logd_7_4"]
NUMERIC_COLUMNS = {"vdss_l_per_kg", "clearance_ml_per_min_per_kg", "fraction_unbound_plasma", "mrt_h", "half_life_h",
                   "year_of_first_disclosure", "mw", "hba", "hbd", "tpsa_no", "rot_bond_count", "moka_logp",
                   "moka_logd_7_4"}


def number_text(value: float | int) -> str:
    if isinstance(value, bool):
        raise ValueError("boolean cell in a numeric column")
    if isinstance(value, int):
        return str(value)
    return str(int(value)) if value == int(value) and abs(value) < 1e15 else repr(value)


def main() -> int:
    workbook_path, output_path, export_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(export_path, encoding="utf-8", errors="replace", newline="") as handle:
        export_rows = list(csv.reader(handle))
    export_start = next(index for index, row in enumerate(export_rows) if row and row[0] == "Name")
    if [name.replace("\r", "") for name in export_rows[export_start][:8]] != SOURCE_HEADER[:8]:
        raise SystemExit("unexpected header in the export")
    export_by_cas = {row[1].strip(): row for row in export_rows[export_start + 1:] if row and row[0]}
    sheet = openpyxl.load_workbook(workbook_path, data_only=False).worksheets[0]
    header = [cell.value for cell in sheet[HEADER_ROW]]
    if header != SOURCE_HEADER:
        raise SystemExit(f"unexpected header: {header}")
    rows_written = 0
    with open(output_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(OUTPUT_HEADER + ["source_row", "note"])
        for row in sheet.iter_rows(min_row=HEADER_ROW + 1, max_row=sheet.max_row):
            values = [cell.value for cell in row[: len(header)]]
            if all(value is None or value == "" for value in values):
                continue
            output_values, notes = [], []
            for column, source_column, value in zip(OUTPUT_HEADER, SOURCE_HEADER, values):
                if value is None or (isinstance(value, str) and value == ""):
                    output_values.append("")
                elif column in NUMERIC_COLUMNS:
                    if isinstance(value, (int, float)):
                        text = number_text(value)
                        output_values.append(text)
                        if len(text.replace("-", "").replace(".", "").lstrip("0")) > 8:
                            shown = export_by_cas[str(values[1]).strip()][SOURCE_HEADER.index(source_column)].strip()
                            notes.append(f"{source_column!r} is stored in the workbook as the long decimal {text}, written here "
                                         f"unrounded; the comma-separated export of the workbook in the same repository shows {shown}")
                    else:
                        output_values.append("")
                        notes.append(f"{source_column!r} cell holds the text {value!r}, left empty here")
                elif isinstance(value, str):
                    if column in ("name", "cas_number") and value != value.strip():
                        notes.append(f"{source_column!r} cell has outer whitespace in the source ({value!r}), stripped here")
                        value = value.strip()
                    output_values.append(value)
                else:
                    raise SystemExit(f"row {row[0].row}: non-text value {value!r} in text column {column}")
            writer.writerow(output_values + [row[0].row, "; ".join(notes)])
            rows_written += 1
    print("rows written:", rows_written)
    return 0


if __name__ == "__main__":
    sys.exit(main())
